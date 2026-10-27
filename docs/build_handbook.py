"""Build ARTECH Studio OS staff handbooks as Word documents."""

from __future__ import annotations

import re
from pathlib import Path

from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_LINE_SPACING, WD_TAB_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

LATIN_CHUNK = re.compile(
    r"(?:[A-Za-z][A-Za-z0-9+./&_\-]*|[0-9]+(?:[.,][0-9]+)?)(?:[ \u00A0]+(?:[A-Za-z][A-Za-z0-9+./&_\-]*|[0-9]+(?:[.,][0-9]+)?))*"
)

OUT = Path(__file__).resolve().parent
INK = RGBColor(0x1A, 0x1A, 0x1A)
MUTED = RGBColor(0x5C, 0x5C, 0x5C)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)
GRAPHITE = "121212"
ROW_ALT = "F4F4F4"
RULE = "D4D4D4"
BANNER = "1C1C1C"


def set_run_font(run, name: str, size: Pt | None = None, bold=False, color=INK, cs_name: str | None = None):
    run.bold = bold
    run.font.color.rgb = color
    if size:
        run.font.size = size
    run.font.name = name
    rPr = run._element.get_or_add_rPr()
    rFonts = rPr.find(qn("w:rFonts"))
    if rFonts is None:
        rFonts = OxmlElement("w:rFonts")
        rPr.append(rFonts)
    rFonts.set(qn("w:ascii"), name)
    rFonts.set(qn("w:hAnsi"), name)
    rFonts.set(qn("w:eastAsia"), name)
    rFonts.set(qn("w:cs"), cs_name or name)


def set_run_rtl(run, on: bool):
    rPr = run._element.get_or_add_rPr()
    existing = rPr.find(qn("w:rtl"))
    if existing is not None:
        rPr.remove(existing)
    rtl = OxmlElement("w:rtl")
    rtl.set(qn("w:val"), "1" if on else "0")
    rPr.append(rtl)
    cs = rPr.find(qn("w:cs"))
    if on:
        if cs is None:
            cs = OxmlElement("w:cs")
            rPr.append(cs)
        cs.set(qn("w:val"), "1")


def add_mixed_text(paragraph, text, *, rtl, bold=False, size=11, color=INK, font=None):
    """Keep English product names LTR inside Arabic (RTL) paragraphs."""
    if not rtl:
        run = paragraph.add_run(text)
        set_run_font(run, font or "Calibri", Pt(size), bold=bold, color=color)
        return
    font = font or "Segoe UI"
    marker = paragraph.add_run("\u200F")
    set_run_font(marker, font, Pt(1), bold=bold, color=color, cs_name="Segoe UI")
    set_run_rtl(marker, True)
    idx = 0
    for match in LATIN_CHUNK.finditer(text):
        if match.start() > idx:
            run = paragraph.add_run(text[idx : match.start()])
            set_run_font(run, font, Pt(size), bold=bold, color=color, cs_name="Segoe UI")
            set_run_rtl(run, True)
        run = paragraph.add_run(match.group())
        set_run_font(run, font, Pt(size), bold=bold, color=color, cs_name="Segoe UI")
        set_run_rtl(run, False)
        idx = match.end()
    if idx < len(text):
        run = paragraph.add_run(text[idx:])
        set_run_font(run, font, Pt(size), bold=bold, color=color, cs_name="Segoe UI")
        set_run_rtl(run, True)


def shade_cell(cell, fill: str):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:fill"), fill)
    shd.set(qn("w:val"), "clear")
    tcPr.append(shd)


def set_cell_borders(cell, color=RULE, sz="4"):
    tcPr = cell._tc.get_or_add_tcPr()
    tcBorders = OxmlElement("w:tcBorders")
    for edge in ("top", "left", "bottom", "right"):
        el = OxmlElement(f"w:{edge}")
        el.set(qn("w:val"), "single")
        el.set(qn("w:sz"), sz)
        el.set(qn("w:space"), "0")
        el.set(qn("w:color"), color)
        tcBorders.append(el)
    tcPr.append(tcBorders)


def set_cell_margins(cell, top=60, bottom=60, left=80, right=80):
    tcPr = cell._tc.get_or_add_tcPr()
    tcMar = OxmlElement("w:tcMar")
    for name, val in (("top", top), ("left", left), ("bottom", bottom), ("right", right)):
        el = OxmlElement(f"w:{name}")
        el.set(qn("w:w"), str(val))
        el.set(qn("w:type"), "dxa")
        tcMar.append(el)
    tcPr.append(tcMar)


def prevent_row_split(row):
    trPr = row._tr.get_or_add_trPr()
    cant = OxmlElement("w:cantSplit")
    trPr.append(cant)


def mark_header_row(row):
    trPr = row._tr.get_or_add_trPr()
    hdr = OxmlElement("w:tblHeader")
    trPr.append(hdr)


def set_bidi(paragraph, on=True):
    pPr = paragraph._p.get_or_add_pPr()
    existing = pPr.find(qn("w:bidi"))
    if existing is not None:
        pPr.remove(existing)
    bidi = OxmlElement("w:bidi")
    bidi.set(qn("w:val"), "1" if on else "0")
    pPr.append(bidi)
    jc = pPr.find(qn("w:jc"))
    if jc is None:
        jc = OxmlElement("w:jc")
        pPr.append(jc)
    jc.set(qn("w:val"), "right" if on else "left")
    paragraph.alignment = WD_ALIGN_PARAGRAPH.RIGHT if on else WD_ALIGN_PARAGRAPH.LEFT


def set_rtl_section(section, on=True):
    sectPr = section._sectPr
    existing = sectPr.find(qn("w:bidi"))
    if existing is not None:
        sectPr.remove(existing)
    if on:
        bidi = OxmlElement("w:bidi")
        bidi.set(qn("w:val"), "1")
        sectPr.append(bidi)


def add_page_number(paragraph, rtl=False):
    font = "Segoe UI" if rtl else "Calibri"
    begin = paragraph.add_run()
    start = OxmlElement("w:fldChar")
    start.set(qn("w:fldCharType"), "begin")
    begin._r.append(start)
    instr_run = paragraph.add_run()
    instr = OxmlElement("w:instrText")
    instr.set(qn("xml:space"), "preserve")
    instr.text = " PAGE "
    instr_run._r.append(instr)
    sep_run = paragraph.add_run()
    sep = OxmlElement("w:fldChar")
    sep.set(qn("w:fldCharType"), "separate")
    sep_run._r.append(sep)
    vis = paragraph.add_run("1")
    set_run_font(vis, font, Pt(9), color=MUTED, cs_name="Segoe UI")
    if rtl:
        set_run_rtl(vis, False)
    end_run = paragraph.add_run()
    end = OxmlElement("w:fldChar")
    end.set(qn("w:fldCharType"), "end")
    end_run._r.append(end)


def add_toc_field(paragraph):
    run = paragraph.add_run()
    r = run._r
    begin = OxmlElement("w:fldChar")
    begin.set(qn("w:fldCharType"), "begin")
    begin.set(qn("w:dirty"), "true")
    r.append(begin)
    instr = OxmlElement("w:instrText")
    instr.set(qn("xml:space"), "preserve")
    instr.text = r' TOC \o "1-1" \h \z \u '
    r.append(instr)
    sep = OxmlElement("w:fldChar")
    sep.set(qn("w:fldCharType"), "separate")
    r.append(sep)
    end_run = paragraph.add_run()
    end = OxmlElement("w:fldChar")
    end.set(qn("w:fldCharType"), "end")
    end_run._r.append(end)


def configure_styles(doc, rtl=False):
    body = "Segoe UI" if rtl else "Calibri"
    heading_font = "Segoe UI"
    styles = doc.styles

    normal = styles["Normal"]
    normal.font.name = body
    normal.font.size = Pt(11)
    normal.font.color.rgb = INK
    pf = normal.paragraph_format
    pf.space_after = Pt(8)
    pf.space_before = Pt(0)
    pf.line_spacing_rule = WD_LINE_SPACING.MULTIPLE
    pf.line_spacing = 1.15
    rPr = normal.element.get_or_add_rPr()
    rFonts = rPr.find(qn("w:rFonts"))
    if rFonts is None:
        rFonts = OxmlElement("w:rFonts")
        rPr.insert(0, rFonts)
    rFonts.set(qn("w:ascii"), body)
    rFonts.set(qn("w:hAnsi"), body)
    rFonts.set(qn("w:cs"), "Segoe UI")
    rFonts.set(qn("w:eastAsia"), body)

    for style_name, size, before, after, bold in (
        ("Title", 28, 0, 6, True),
        ("Subtitle", 16, 0, 12, False),
        ("Heading 1", 18, 22, 8, True),
        ("Heading 2", 13, 16, 6, True),
        ("Heading 3", 12, 12, 4, True),
    ):
        st = styles[style_name]
        st.font.name = heading_font
        st.font.size = Pt(size)
        st.font.bold = bold
        st.font.color.rgb = INK
        st.paragraph_format.space_before = Pt(before)
        st.paragraph_format.space_after = Pt(after)
        st.paragraph_format.keep_with_next = True
        st.paragraph_format.line_spacing = 1.15
        rPr = st.element.get_or_add_rPr()
        rFonts = rPr.find(qn("w:rFonts"))
        if rFonts is None:
            rFonts = OxmlElement("w:rFonts")
            rPr.insert(0, rFonts)
        rFonts.set(qn("w:ascii"), heading_font)
        rFonts.set(qn("w:hAnsi"), heading_font)
        rFonts.set(qn("w:cs"), "Segoe UI")

    h1 = styles["Heading 1"].element.get_or_add_pPr()
    pbdr = OxmlElement("w:pBdr")
    bottom = OxmlElement("w:bottom")
    bottom.set(qn("w:val"), "single")
    bottom.set(qn("w:sz"), "12")
    bottom.set(qn("w:space"), "4")
    bottom.set(qn("w:color"), GRAPHITE)
    pbdr.append(bottom)
    h1.append(pbdr)


def setup_page(doc, rtl=False):
    section = doc.sections[0]
    section.page_width = Cm(21.0)
    section.page_height = Cm(29.7)
    section.left_margin = Cm(2.0)
    section.right_margin = Cm(2.0)
    section.top_margin = Cm(2.0)
    section.bottom_margin = Cm(2.2)
    section.header_distance = Cm(1.0)
    section.footer_distance = Cm(1.0)
    section.different_first_page_header_footer = True
    set_rtl_section(section, rtl)

    settings = doc.settings.element
    lang = OxmlElement("w:themeFontLang")
    lang.set(qn("w:val"), "ar-EG" if rtl else "en-GB")
    lang.set(qn("w:bidi"), "ar-EG")
    settings.append(lang)

    def fill_header(hp, text):
        hp.clear()
        hp.paragraph_format.space_after = Pt(4)
        if rtl:
            set_bidi(hp, True)
            hp.alignment = WD_ALIGN_PARAGRAPH.RIGHT
            add_mixed_text(hp, text, rtl=True, size=9, color=MUTED)
        else:
            hp.alignment = WD_ALIGN_PARAGRAPH.LEFT
            add_mixed_text(hp, text, rtl=False, size=9, color=MUTED, font="Calibri")

    header = section.header
    header.is_linked_to_previous = False
    fill_header(
        header.paragraphs[0],
        "دليل الموظفين  ·  ARTECH Studio OS" if rtl else "ARTECH Studio OS  ·  Staff Handbook",
    )
    first_header = section.first_page_header
    first_header.is_linked_to_previous = False
    first_header.paragraphs[0].clear()

    footer = section.footer
    footer.is_linked_to_previous = False
    fp = footer.paragraphs[0]
    fp.clear()
    if rtl:
        set_bidi(fp, True)
        fp.alignment = WD_ALIGN_PARAGRAPH.CENTER
        add_page_number(fp, rtl=True)
        sep = fp.add_run("  ·  ")
        set_run_font(sep, "Segoe UI", Pt(9), color=MUTED, cs_name="Segoe UI")
        add_mixed_text(fp, "للاستخدام الداخلي", rtl=True, size=9, color=MUTED)
    else:
        fp.alignment = WD_ALIGN_PARAGRAPH.LEFT
        fp.paragraph_format.tab_stops.add_tab_stop(Cm(17.0), WD_TAB_ALIGNMENT.RIGHT)
        label = fp.add_run("Internal use")
        set_run_font(label, "Calibri", Pt(9), color=MUTED)
        fp.add_run("\t")
        add_page_number(fp, rtl=False)
    section.first_page_footer.is_linked_to_previous = False
    section.first_page_footer.paragraphs[0].clear()


def usable_width(doc):
    s = doc.sections[0]
    return s.page_width - s.left_margin - s.right_margin


def p(doc, text, *, rtl=False, size=11, bold=False, color=INK, space_after=8, space_before=0, align=None, font=None):
    para = doc.add_paragraph()
    para.paragraph_format.space_after = Pt(space_after)
    para.paragraph_format.space_before = Pt(space_before)
    para.paragraph_format.line_spacing = 1.15
    if rtl:
        set_bidi(para, True)
        para.alignment = align or WD_ALIGN_PARAGRAPH.RIGHT
    else:
        para.alignment = align or WD_ALIGN_PARAGRAPH.LEFT
    add_mixed_text(para, text, rtl=rtl, bold=bold, size=size, color=color, font=font)
    return para


def rich(doc, parts, *, rtl=False, space_after=8, size=11):
    """parts: list of str | (text, bold)."""
    para = doc.add_paragraph()
    para.paragraph_format.space_after = Pt(space_after)
    para.paragraph_format.line_spacing = 1.15
    if rtl:
        set_bidi(para, True)
        para.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    else:
        para.alignment = WD_ALIGN_PARAGRAPH.LEFT
    for part in parts:
        text, bold = part if isinstance(part, tuple) else (part, False)
        add_mixed_text(para, text, rtl=rtl, bold=bold, size=size)
    return para


AR_DIGITS = str.maketrans("0123456789", "٠١٢٣٤٥٦٧٨٩")


def eastern_prefix(text: str) -> str:
    return re.sub(r"^(\d+)\.\s*", lambda m: m.group(1).translate(AR_DIGITS) + ".  ", text)


def heading(doc, text, level=1, rtl=False):
    h = doc.add_heading("", level=level)
    h.paragraph_format.keep_with_next = True
    size = 18 if level == 1 else 13
    if rtl:
        set_bidi(h, True)
        h.alignment = WD_ALIGN_PARAGRAPH.RIGHT
        text = eastern_prefix(text)
    add_mixed_text(h, text, rtl=rtl, bold=True, size=size, font="Segoe UI")
    return h


def write_parts(para, item, *, rtl):
    parts = item if isinstance(item, list) else [item]
    for part in parts:
        text, bold = part if isinstance(part, tuple) else (part, False)
        add_mixed_text(para, text, rtl=rtl, bold=bold, size=11)


def bullets(doc, items, *, rtl=False):
    for item in items:
        para = doc.add_paragraph()
        para.paragraph_format.space_after = Pt(4)
        para.paragraph_format.line_spacing = 1.15
        if rtl:
            set_bidi(para, True)
            para.alignment = WD_ALIGN_PARAGRAPH.RIGHT
            para.paragraph_format.right_indent = Cm(0.4)
            add_mixed_text(para, "•  ", rtl=True, size=11)
            write_parts(para, item, rtl=True)
        else:
            para.paragraph_format.left_indent = Cm(0.75)
            para.paragraph_format.first_line_indent = Cm(-0.4)
            add_mixed_text(para, "•  ", rtl=False, size=11)
            write_parts(para, item, rtl=False)


def numbered(doc, items, *, rtl=False):
    for i, item in enumerate(items, 1):
        para = doc.add_paragraph()
        para.paragraph_format.space_after = Pt(6)
        para.paragraph_format.line_spacing = 1.15
        prefix = f"{str(i).translate(AR_DIGITS)}.  " if rtl else f"{i}.  "
        if rtl:
            set_bidi(para, True)
            para.alignment = WD_ALIGN_PARAGRAPH.RIGHT
            para.paragraph_format.right_indent = Cm(0.4)
            add_mixed_text(para, prefix, rtl=True, bold=True, size=11)
            write_parts(para, item, rtl=True)
        else:
            para.paragraph_format.left_indent = Cm(1.0)
            para.paragraph_format.first_line_indent = Cm(-0.75)
            add_mixed_text(para, prefix, rtl=False, bold=True, size=11)
            write_parts(para, item, rtl=False)


def note_box(doc, text, *, rtl=False):
    table = doc.add_table(rows=1, cols=1)
    table.autofit = True
    table.allow_autofit = True
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    cell = table.cell(0, 0)
    shade_cell(cell, "F3F3F3")
    set_cell_margins(cell, top=80, bottom=80, left=120, right=120)
    tcPr = cell._tc.get_or_add_tcPr()
    tcBorders = OxmlElement("w:tcBorders")
    for edge, sz, color in (
        ("top", "4", RULE),
        ("bottom", "4", RULE),
        ("left", "24" if not rtl else "4", GRAPHITE if not rtl else RULE),
        ("right", "24" if rtl else "4", GRAPHITE if rtl else RULE),
    ):
        el = OxmlElement(f"w:{edge}")
        el.set(qn("w:val"), "single")
        el.set(qn("w:sz"), sz)
        el.set(qn("w:space"), "0")
        el.set(qn("w:color"), color)
        tcBorders.append(el)
    tcPr.append(tcBorders)
    para = cell.paragraphs[0]
    para.paragraph_format.space_after = Pt(0)
    if rtl:
        set_bidi(para, True)
        para.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    add_mixed_text(para, text, rtl=rtl, size=10.5)
    doc.add_paragraph().paragraph_format.space_after = Pt(8)


def add_table(doc, headers, rows, *, rtl=False, col_widths=None):
    table = doc.add_table(rows=1 + len(rows), cols=len(headers))
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = True
    if rtl:
        bidi_visual = OxmlElement("w:bidiVisual")
        bidi_visual.set(qn("w:val"), "1")
        table._tbl.tblPr.append(bidi_visual)
    width = usable_width(doc)
    if col_widths:
        total = sum(col_widths)
        widths = [int(width * w / total) for w in col_widths]
    else:
        widths = [int(width / len(headers))] * len(headers)

    tbl = table._tbl
    tblPr = tbl.tblPr if tbl.tblPr is not None else OxmlElement("w:tblPr")
    tblW = OxmlElement("w:tblW")
    tblW.set(qn("w:w"), str(int(width.twips if hasattr(width, "twips") else width / 635)))
    # python-docx Length is EMU; 914400 EMU = 1 inch = 1440 twips. EMU/635 ≈ twips.
    tblW.set(qn("w:w"), str(int(width / 635)))
    tblW.set(qn("w:type"), "dxa")
    tblPr.append(tblW)

    font = "Segoe UI" if rtl else "Calibri"
    header_row = table.rows[0]
    mark_header_row(header_row)
    prevent_row_split(header_row)
    for i, label in enumerate(headers):
        cell = header_row.cells[i]
        shade_cell(cell, BANNER)
        set_cell_borders(cell, color="1C1C1C", sz="4")
        set_cell_margins(cell)
        cell.width = widths[i]
        para = cell.paragraphs[0]
        para.paragraph_format.space_after = Pt(0)
        para.paragraph_format.space_before = Pt(0)
        if rtl:
            set_bidi(para, True)
            para.alignment = WD_ALIGN_PARAGRAPH.RIGHT
        add_mixed_text(para, label, rtl=rtl, bold=True, size=10, color=WHITE, font=font)

    for r_i, row_data in enumerate(rows):
        row = table.rows[r_i + 1]
        prevent_row_split(row)
        fill = ROW_ALT if r_i % 2 else "FFFFFF"
        for c_i, value in enumerate(row_data):
            cell = row.cells[c_i]
            shade_cell(cell, fill)
            set_cell_borders(cell, color=RULE, sz="4")
            set_cell_margins(cell)
            cell.width = widths[c_i]
            para = cell.paragraphs[0]
            para.paragraph_format.space_after = Pt(0)
            para.paragraph_format.space_before = Pt(0)
            if rtl:
                set_bidi(para, True)
                para.alignment = WD_ALIGN_PARAGRAPH.RIGHT
            add_mixed_text(para, value, rtl=rtl, size=10, color=INK, font=font)

    spacer = doc.add_paragraph()
    spacer.paragraph_format.space_after = Pt(10)
    return table


def cover(doc, *, rtl=False):
    table = doc.add_table(rows=1, cols=1)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    cell = table.cell(0, 0)
    shade_cell(cell, GRAPHITE)
    set_cell_margins(cell, top=360, bottom=360, left=240, right=240)
    tcPr = cell._tc.get_or_add_tcPr()
    tcBorders = OxmlElement("w:tcBorders")
    for edge in ("top", "left", "bottom", "right"):
        el = OxmlElement(f"w:{edge}")
        el.set(qn("w:val"), "nil")
        tcBorders.append(el)
    tcPr.append(tcBorders)

    def banner_p(text, size, bold=False, after=4, before=0, color=WHITE):
        para = cell.add_paragraph()
        para.paragraph_format.space_after = Pt(after)
        para.paragraph_format.space_before = Pt(before)
        if rtl:
            set_bidi(para, True)
            para.alignment = WD_ALIGN_PARAGRAPH.RIGHT
        else:
            para.alignment = WD_ALIGN_PARAGRAPH.LEFT
        add_mixed_text(para, text, rtl=rtl, bold=bold, size=size, color=color, font="Segoe UI")
        return para

    first = cell.paragraphs[0]
    first.paragraph_format.space_after = Pt(6)
    if rtl:
        set_bidi(first, True)
        first.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    add_mixed_text(first, "ARTECH", rtl=rtl, bold=True, size=14, color=RGBColor(0xC8, 0xC8, 0xC8), font="Segoe UI")
    banner_p("Studio OS", 28, True, after=8)
    if rtl:
        banner_p("دليل الموظفين", 20, False, after=16, color=RGBColor(0xE8, 0xE8, 0xE8))
        banner_p("كيف تستخدم النظام يومًا بيوم — وليس دليل تركيب تقني", 12, False, after=4, color=RGBColor(0xC8, 0xC8, 0xC8))
        banner_p("للاستخدام الداخلي  ·  أكتوبر ٢٠٢٦", 11, False, after=0, color=RGBColor(0xA8, 0xA8, 0xA8))
    else:
        banner_p("Staff handbook", 20, False, after=16, color=RGBColor(0xE8, 0xE8, 0xE8))
        banner_p("How to use the system day to day — not a software installation guide.", 12, False, after=4, color=RGBColor(0xC8, 0xC8, 0xC8))
        banner_p("October 2026  ·  Internal use", 11, False, after=0, color=RGBColor(0xA8, 0xA8, 0xA8))

    intro = (
        "للموظفين في كل الأقسام: الأعمال، التطوير، العمليات، والحسابات. يشرح الأدوار، مسار المشروع، وكل شاشة في النظام."
        if rtl
        else "For everyone in Business, Development, Operations, and Accounting. It explains roles, the project path, and every screen in the system."
    )
    p(doc, intro, rtl=rtl, size=12, color=MUTED, space_before=16, space_after=0)
    doc.add_page_break()


def toc_page(doc, *, rtl=False):
    title = doc.add_paragraph()
    title.paragraph_format.space_after = Pt(12)
    title.paragraph_format.space_before = Pt(0)
    if rtl:
        set_bidi(title, True)
        title.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    add_mixed_text(title, "المحتويات" if rtl else "Contents", rtl=rtl, bold=True, size=18, font="Segoe UI")
    pPr = title._p.get_or_add_pPr()
    outline = OxmlElement("w:outlineLvl")
    outline.set(qn("w:val"), "9")
    pPr.append(outline)
    pbdr = OxmlElement("w:pBdr")
    bottom = OxmlElement("w:bottom")
    bottom.set(qn("w:val"), "single")
    bottom.set(qn("w:sz"), "12")
    bottom.set(qn("w:space"), "4")
    bottom.set(qn("w:color"), GRAPHITE)
    pbdr.append(bottom)
    pPr.append(pbdr)

    para = doc.add_paragraph()
    if rtl:
        set_bidi(para, True)
        para.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    add_toc_field(para)
    doc.add_page_break()


def build_english(path: Path):
    doc = Document()
    configure_styles(doc, rtl=False)
    setup_page(doc, rtl=False)
    doc.core_properties.author = "ARTECH"
    doc.core_properties.title = "ARTECH Studio OS — Staff Handbook (English)"
    doc.core_properties.subject = "User guide for Studio OS"
    doc.core_properties.category = "Handbook"

    cover(doc, rtl=False)
    toc_page(doc, rtl=False)

    heading(doc, "1.  What Studio OS is")
    p(
        doc,
        "Studio OS is ARTECH’s system for running interactive and experiential projects in one place. Postgres is the source of truth. What you see depends on your department and whether you are a head.",
    )
    p(doc, "The system covers:", bold=True, space_after=4)
    bullets(
        doc,
        [
            "Client briefs",
            "Developer assignment and deadlines",
            "Client feedback",
            "Internal requests (blocked / needs data / needs assets)",
            "Equipment",
            "Site installation",
            "Budget, collections, expenses, and profit",
            "Files shared with the team",
            "Email reminders",
        ],
    )
    note_box(
        doc,
        "Switch the app language with العربية / English in the header. Dates and money follow Cairo time.",
    )

    heading(doc, "2.  Sign in and your account")
    numbered(
        doc,
        [
            "Open the Studio OS address your company uses.",
            [
                "Sign in with the email and temporary password your department head (or ",
                ("Super Admin", True),
                ") gave you.",
            ],
            [
                "On first login you ",
                ("must set a new password", True),
                " on ",
                ("Account", True),
                " before you can use the rest of the app.",
            ],
            ["Later, change your password any time from ", ("Account", True), "."],
        ],
    )
    p(doc, "If an account is disabled, sign-in is blocked. Ask your department head.")
    p(
        doc,
        "The microphone on some fields is dictation (speech to text). Arabic is the default so names and notes are captured correctly.",
    )

    heading(doc, "3.  Roles")
    add_table(
        doc,
        ["Role", "Department", "Typical job in the system"],
        [
            ["Super Admin", "—", "Sees everything. Appoints heads. Users, activity log, automations."],
            [
                "Business Head",
                "Business",
                "Briefs, budget, installation date, client feedback decisions, mark completed. Manages Business users.",
            ],
            ["Business", "Business", "Same project work as Business Head, without managing users or exporting Excel."],
            ["Accountant Manager", "Accounting", "Budget, payments, expenses. Manages Accounting users."],
            ["Accountant", "Accounting", "Same finance work, without managing users."],
            [
                "Tech Head",
                "Development",
                "Assigns developers, deadlines, milestones, reviews work. Sees all projects and budget (view). Manages developers.",
            ],
            [
                "Developer",
                "Development",
                "Only projects they are assigned to. Work, milestones, installation guide, equipment requests.",
            ],
            [
                "Operations Manager",
                "Operations",
                "Installation calendar, start/finish install, equipment, suppliers. Manages Operations users.",
            ],
            ["Operations", "Operations", "Same site work, without managing users."],
        ],
        col_widths=[2.2, 1.8, 5.0],
    )
    p(
        doc,
        "Heads can add, edit, reset password, and disable regular members of their own department. They cannot change other heads or the Super Admin. Only Super Admin appoints heads.",
    )

    heading(doc, "4.  Who sees which menu")
    add_table(
        doc,
        ["Menu", "Who sees it"],
        [
            ["Home, Projects, Requests, Account", "Everyone"],
            ["Assignments", "Super Admin, Tech Head"],
            ["Installations", "Everyone except Developers"],
            ["Equipment", "Super Admin, Development, Operations"],
            ["Suppliers", "Super Admin, Business, Operations"],
            ["Finance", "Super Admin, Business, Accounting, Tech Head"],
            ["Users", "Super Admin and department heads"],
            ["Activity, Automations", "Super Admin only"],
        ],
        col_widths=[4, 6],
    )
    note_box(doc, "Developers never see the real installation date or any money.")

    heading(doc, "5.  The life of a project")
    p(doc, "This is the normal path. Buttons only appear when your role can take that step.")
    add_table(
        doc,
        ["Step", "Status", "Who acts", "What they do"],
        [
            ["1", "Brief submitted", "Tech Head", "Assigns the team and development deadline."],
            ["2", "Assigned", "Developer", "Starts development."],
            ["3", "In development", "Developer", "Sends the job for internal review."],
            ["4", "Internal review", "Tech Head", "Approves only if the guide and demo video exist."],
            ["5", "Ready for installation", "Operations", "Starts the site install."],
            ["6", "Installing", "Operations", "Marks the job installed and writes the report."],
            ["7", "Installed", "Business", "Marks the job completed."],
            ["8", "Completed", "—", "The monthly report uses this date."],
        ],
        col_widths=[1.0, 2.6, 2.0, 4.4],
    )
    bullets(
        doc,
        [
            ["Tech Head can send a project ", ("back to development", True), " from internal review or from ready for installation."],
            ["Operations can pause an install (back to ready for installation)."],
            "A project can also be On hold or Cancelled in the data; day-to-day work uses the path above.",
        ],
    )
    heading(doc, "Who does what on that path", 2)
    numbered(
        doc,
        [
            [
                ("Business", True),
                " creates the brief (name, client, description, references, design, content, other notes). Sets budget and the ",
                ("real installation date", True),
                " when known. Can add progress-showing dates (when WIP is shown to the client).",
            ],
            [
                ("Tech Head", True),
                " opens ",
                ("Assignments", True),
                ", picks developers, sets the development deadline and milestones.",
            ],
            [
                ("Developers", True),
                " start development, tick milestones, write the ",
                ("installation guide", True),
                " (English and/or Arabic) and add a ",
                ("demo video", True),
                " link recorded on the same setup that will go on site.",
            ],
            "Tech Head cannot approve “ready for installation” until steps and demo video are there.",
            [
                ("Operations", True),
                " uses ",
                ("Installations", True),
                ", assigns a lead technician, starts the job, writes the report, marks installed.",
            ],
            [
                ("Business", True),
                " marks the job ",
                ("completed", True),
                " after install (this date is used in the monthly report).",
            ],
            [
                ("Accounting", True),
                " records collections and expenses at any time after the budget exists.",
            ],
        ],
    )

    heading(doc, "6.  Opening a project — tabs")
    p(doc, "A project is split into tabs. You can always open any tab you are allowed to see.")
    p(doc, "If you open a project from Home or Projects, the first tab depends on your role:", bold=True, space_after=6)
    add_table(
        doc,
        ["You are", "Lands on"],
        [
            ["Accountant", "Finance"],
            ["Operations", "Installation"],
            ["Developer / Tech Head", "Development"],
            ["Business / Super Admin", "Brief"],
        ],
        col_widths=[5, 5],
    )
    p(doc, "If you click a project from another page, that page wins:", bold=True, space_after=6)
    add_table(
        doc,
        ["You clicked from", "Opens"],
        [
            ["Finance", "Finance"],
            ["Installations", "Installation"],
            ["Assignments", "Development"],
            ["Requests", "Requests"],
            ["Equipment", "Equipment"],
        ],
        col_widths=[5, 5],
    )

    heading(doc, "7.  Home")
    p(
        doc,
        "Greeting, your role, and (if you can see all projects) pipeline counts and upcoming development deadlines. Widgets appear only when they have something for you:",
    )
    bullets(
        doc,
        [
            [("My projects", True), " — jobs you are assigned to"],
            [("Upcoming installations", True), " — if you handle operations"],
            [("Client feedback", True), " assigned to you, or waiting for Business"],
            [("System status", True), " — Super Admin only"],
        ],
    )

    heading(doc, "8.  Projects")
    bullets(
        doc,
        [
            [("Developers", True), " see only assigned projects."],
            "Everyone else sees all projects (except cancelled ones in some lists).",
            [("New brief", True), " is for Business (and Super Admin)."],
            [
                ("Download Excel", True),
                " is for Super Admin and heads of Business or Accounting. You can group quarters by brief date or by installation date.",
            ],
        ],
    )
    p(doc, "Each card shows status, client, team, development deadline, and (if you are allowed) installation date and budget.")

    heading(doc, "9.  Project tabs in detail")

    heading(doc, "Brief", 2)
    p(doc, "What the client asked for. Business can edit. Progress dates (when work in progress is shown) live here too.")

    heading(doc, "Feedback", 2)
    p(
        doc,
        "Log what the client said (WhatsApp, call, email, meeting, site visit, or the client link). Split into items: design, content, technical, new request, general.",
    )
    p(doc, "Typical path:", bold=True, space_after=4)
    numbered(
        doc,
        [
            ["Business or Operations ", ("logs", True), " feedback (optionally with AI split into items)."],
            [
                "If it is a ",
                ("new request not in the brief", True),
                ", it waits for Business to decide: included / extra cost / rejected.",
            ],
            ["Tech Head or Business ", ("assigns", True), " an item."],
            ["The assignee (or developers on the project) mark ", ("in progress", True), ", then ", ("done", True), "."],
            [
                "The developer can ask for ",
                ("extra days", True),
                " on the development deadline; Tech Head accepts or declines.",
            ],
            [
                "Business confirms the ",
                ("client approved", True),
                ", or the client does that on their own link.",
            ],
        ],
    )
    note_box(
        doc,
        "Client link: Business can turn on a private link. The client sees a simple status (not internal steps), can add comments, and can say an item is not right yet. Team discussion stays internal. Urgent items are marked clearly.",
    )

    heading(doc, "Files", 2)
    p(
        doc,
        "Shared uploads (cloud): Brief, Client assets, Brand, Design, Content, Preview, Delivery. Maximum 100 MB per file.",
    )
    p(
        doc,
        "Heavy working files (Unreal, C4D, Premiere) stay on the office disk folders shown on the page — they are not uploaded.",
    )
    p(doc, "Developers can manage files only on projects they are on. Business, Operations, Tech Head, and Super Admin can manage files more broadly.")

    heading(doc, "Development", 2)
    p(doc, "Team, development deadline, milestones, status buttons, and the installation guide.")
    bullets(
        doc,
        [
            [("Tech Head:", True), " assign / remove developers, deadline, add milestones, approve or send back."],
            [
                ("Developers on the project:", True),
                " start work, send for review, tick milestones, write the guide (either language; the other can be translated on save if empty).",
            ],
        ],
    )

    heading(doc, "Requests", 2)
    p(
        doc,
        "Inside a project and on the Requests page. Ask a person for what this job needs: blocked, needs data, needs assets, or other. Only people involved see each request.",
    )
    bullets(
        doc,
        [
            ["Recipients (or their department head) resolve ", ("their part", True), "."],
            "When every part is resolved, the request is closed.",
            ["The person who raised it (or Super Admin) can ", ("cancel", True), "."],
            ["A badge on ", ("Requests", True), " in the menu is the count waiting on you."],
        ],
    )

    heading(doc, "Equipment (on the project)", 2)
    p(
        doc,
        "Request units for a date range. If enough units are free, the request is approved immediately. If not, it is waitlisted until stock is free.",
    )
    p(
        doc,
        "When the job is done with the kit, mark Returned (or Cancel). Super Admin and heads of Development or Operations maintain the catalogue (quantity, status: active / maintenance / broken / retired).",
    )

    heading(doc, "Installation", 2)
    p(doc, "Hidden from developers.")
    bullets(
        doc,
        [
            [("Business", True), " sets the real date, location link, and notes."],
            [("Operations", True), " sets lead technician, starts install, pauses, marks installed, writes the report."],
            [("Business", True), " marks ", ("completed", True), " after installed."],
        ],
    )

    heading(doc, "Finance", 2)
    p(doc, "Hidden from developers and Operations.")
    add_table(
        doc,
        ["Figure", "Meaning"],
        [
            ["Budget", "Agreed amount. Does not change when you add expenses."],
            ["Revenue", "Money collected. Starts at 0 until a payment is recorded."],
            ["Outstanding", "Budget minus revenue (still to collect)."],
            ["Expenses", "Costs."],
            ["Profit", "Revenue minus expenses. Shown in red if costs exceed collections."],
        ],
        col_widths=[2.4, 7.6],
    )
    p(doc, "Who can do what:", bold=True, space_after=4)
    bullets(
        doc,
        [
            [("View:", True), " Super Admin, Business, Accounting, Tech Head"],
            [("Edit budget:", True), " Super Admin, Business, Accounting"],
            [("Record payments and expenses:", True), " Super Admin and Accounting only"],
        ],
    )
    note_box(
        doc,
        "How Accounting records many lines: open Record payment or expense → fill a line → Add (repeat) → Save once. You can remove a queued line before save. If the last line is filled and you click Save without Add, that last line is included.",
    )
    p(doc, "The Finance page shows the same figures for every project.")

    heading(doc, "10.  Assignments (Tech Head)")
    p(
        doc,
        "Waiting briefs as cards, then each developer with their active projects underneath (nearest deadline first). Open a card to assign people, deadline, and milestones on the Development tab.",
    )

    heading(doc, "11.  Installations calendar")
    p(
        doc,
        "Month view of real site dates, “installing now”, and projects still without a date. Business sets the date on the project. Operations runs the day on site from the Installation tab.",
    )

    heading(doc, "12.  Equipment catalogue")
    p(
        doc,
        "Stock, free vs in use, which project is holding units, waitlist. Request from a project. Heads of Development/Operations (and Super Admin) add and edit items.",
    )

    heading(doc, "13.  Suppliers")
    p(doc, "Company name, contact, phone, address, and rented equipment. For Business, Operations, and Super Admin.")

    heading(doc, "14.  Users")
    bullets(
        doc,
        [
            "Super Admin sees everyone and can appoint heads.",
            [
                "A head sees only their department and can add members with a ",
                ("temporary password", True),
                " (at least 8 characters). The new person must change it on first login.",
            ],
            [
                "Reset password, disable, or re-enable from the person’s page. You cannot disable yourself; use ",
                ("Account", True),
                " for your own password.",
            ],
        ],
    )

    heading(doc, "15.  Emails the system sends")
    p(doc, "You do not need to trigger these. Typical messages:")
    bullets(
        doc,
        [
            "You were assigned to a project",
            "A request needs your part",
            "A request was cancelled",
            "Client feedback assigned to you",
            "Development deadline in 3 days or tomorrow, or overdue",
            "Installation in 2 days, or overdue",
            "Equipment still in use",
            [
                ("Weekly summary", True),
                " (heads and Super Admin): pipeline, upcoming installs, overdue deadlines, open requests, equipment in use",
            ],
            [
                ("Monthly report", True),
                " (Business Head and Super Admin): completed jobs for the previous month, with a spreadsheet",
            ],
        ],
    )

    heading(doc, "16.  Super Admin only")
    bullets(
        doc,
        [
            [("Activity", True), " — who changed what"],
            [("Automations", True), " — email health, run daily / weekly / monthly jobs if needed"],
        ],
    )

    heading(doc, "17.  Cheat sheets")

    heading(doc, "I am Business", 2)
    numbered(
        doc,
        [
            [("Projects → New brief", True)],
            ["Set ", ("budget", True), " and ", ("installation date", True), " when you have them"],
            ["Log ", ("client feedback", True), "; decide scope on “new request”"],
            ["Turn the ", ("client link", True), " on when the client should reply themselves"],
            ["After Operations marks installed → ", ("Mark completed", True)],
        ],
    )

    heading(doc, "I am Tech Head", 2)
    numbered(
        doc,
        [
            [("Assignments", True), " — pick up waiting briefs"],
            "Assign developers, deadline, milestones",
            [
                "Developers send for review → check guide + demo → ",
                ("Approve: ready for installation", True),
                " (or send back)",
            ],
            "Decide extra days asked from feedback items",
        ],
    )

    heading(doc, "I am a Developer", 2)
    numbered(
        doc,
        [
            [("Home / Projects", True), " — only your jobs"],
            [("Start development", True), " → work milestones → write guide + demo video"],
            [("Send for internal review", True)],
            "Request equipment and raise requests if you are blocked",
            "You will not see money or the real install date",
        ],
    )

    heading(doc, "I am Operations", 2)
    numbered(
        doc,
        [
            [("Installations", True), " calendar"],
            [
                "On the project ",
                ("Installation", True),
                " tab: lead technician, ",
                ("Start installation", True),
                ", report, ",
                ("Mark installed", True),
            ],
            "Equipment and suppliers as needed",
            "You can log client feedback from the site; Business decides extra scope",
        ],
    )

    heading(doc, "I am Accounting", 2)
    numbered(
        doc,
        [
            [("Finance", True), " or a project ", ("Finance", True), " tab (opens automatically)"],
            ["Confirm ", ("budget", True)],
            [("Add", True), " payments and expenses to a list, then ", ("Save", True)],
            ["Watch ", ("revenue", True), " (collections), ", ("outstanding", True), ", ", ("profit", True)],
        ],
    )

    heading(doc, "18.  If something is missing")
    add_table(
        doc,
        ["What you see", "What it usually means"],
        [
            ["A menu item is missing", "Your role is not meant to use it. That is normal."],
            ["You cannot open a project", "You are a developer and not on the team. Ask the Tech Head."],
            ["You cannot change status", "The previous step is not done, or the guide/demo is missing."],
            ["Login fails", "Wrong password, or the account is disabled. Ask your head."],
        ],
        col_widths=[3.5, 6.5],
    )
    p(doc, "Questions about how the company wants a job run (not the buttons) go to your department head.")

    doc.save(path)


def build_arabic(path: Path):
    doc = Document()
    configure_styles(doc, rtl=True)
    setup_page(doc, rtl=True)
    doc.core_properties.author = "ARTECH"
    doc.core_properties.title = "نظام ARTECH Studio OS — دليل الموظفين (العربية)"
    doc.core_properties.subject = "دليل استخدام Studio OS"
    doc.core_properties.category = "Handbook"
    rtl = True

    cover(doc, rtl=True)
    toc_page(doc, rtl=True)

    heading(doc, "1.  ما هو Studio OS", rtl=rtl)
    p(
        doc,
        "نظام ARTECH لإدارة مشاريع التجارب التفاعلية في مكان واحد. قاعدة البيانات هي المرجع الوحيد. ما تراه يعتمد على قسمك وهل أنت رئيس قسم.",
        rtl=rtl,
    )
    p(doc, "يغطي النظام:", rtl=rtl, bold=True, space_after=4)
    bullets(
        doc,
        [
            "ملخصات العملاء",
            "تكليف المطوّرين والمواعيد",
            "ملاحظات العميل",
            "الطلبات الداخلية (متوقف / يحتاج بيانات / يحتاج ملفات)",
            "المعدات",
            "التركيب في الموقع",
            "الميزانية والتحصيل والمصروفات والربح",
            "ملفات الفريق",
            "تذكيرات بالبريد",
        ],
        rtl=rtl,
    )
    note_box(
        doc,
        "بدّل لغة الواجهة من العربية / English في أعلى الصفحة. التواريخ والمبالغ حسب توقيت القاهرة.",
        rtl=rtl,
    )

    heading(doc, "2.  تسجيل الدخول والحساب", rtl=rtl)
    numbered(
        doc,
        [
            "افتح عنوان Studio OS الذي تستخدمه الشركة.",
            [
                "سجّل بالبريد وكلمة المرور المؤقتة التي أعطاها لك رئيس قسمك (أو ",
                ("المدير العام", True),
                ").",
            ],
            [
                "في أول دخول ",
                ("يجب تغيير كلمة المرور", True),
                " من ",
                ("الحساب", True),
                " قبل استخدام باقي النظام.",
            ],
            ["بعدها غيّر كلمة المرور متى شئت من ", ("الحساب", True), "."],
        ],
        rtl=rtl,
    )
    p(doc, "إذا عُطّل الحساب لن تدخل. راجع رئيس قسمك.", rtl=rtl)
    p(
        doc,
        "الميكروفون في بعض الحقول للإملاء الصوتي. العربية هي الافتراضي حتى تُحفظ الأسماء والملاحظات بشكل صحيح.",
        rtl=rtl,
    )

    heading(doc, "3.  الأدوار", rtl=rtl)
    add_table(
        doc,
        ["الدور", "القسم", "العمل المعتاد"],
        [
            ["المدير العام", "—", "يرى كل شيء. يعيّن رؤساء الأقسام. المستخدمون، سجل النشاط، الأتمتة."],
            [
                "رئيس قسم الأعمال",
                "الأعمال",
                "الملخصات، الميزانية، موعد التركيب، قرارات ملاحظات العميل، تعليم المشروع مكتملًا. إدارة مستخدمي الأعمال.",
            ],
            ["أعمال", "الأعمال", "نفس عمل المشروع بدون إدارة المستخدمين أو تصدير Excel."],
            ["مدير الحسابات", "الحسابات", "الميزانية والدفعات والمصروفات. إدارة مستخدمي الحسابات."],
            ["محاسب", "الحسابات", "نفس العمل المالي بدون إدارة المستخدمين."],
            [
                "رئيس القسم التقني",
                "التطوير",
                "تكليف المطوّرين والمواعيد والمراحل ومراجعة العمل. يرى كل المشاريع والميزانية (عرض). إدارة المطوّرين.",
            ],
            ["مطوّر", "التطوير", "المشاريع المكلَّف بها فقط. العمل، المراحل، دليل التركيب، طلب المعدات."],
            [
                "مدير العمليات",
                "العمليات",
                "تقويم التركيب، بدء التركيب وإنهاؤه، المعدات، المورّدون. إدارة مستخدمي العمليات.",
            ],
            ["عمليات", "العمليات", "نفس عمل الموقع بدون إدارة المستخدمين."],
        ],
        rtl=rtl,
        col_widths=[2.2, 1.8, 5.0],
    )
    p(
        doc,
        "رئيس القسم يضيف ويعدّل ويعيد تعيين كلمة المرور ويعطّل أعضاء قسمه العاديين فقط. لا يغيّر رؤساء آخرين ولا المدير العام. المدير العام وحده يعيّن الرؤساء.",
        rtl=rtl,
    )

    heading(doc, "4.  من يرى أي قائمة", rtl=rtl)
    add_table(
        doc,
        ["القائمة", "من يراها"],
        [
            ["الرئيسية، المشاريع، الطلبات، الحساب", "الجميع"],
            ["التكليفات", "المدير العام ورئيس القسم التقني"],
            ["التركيبات", "الجميع ما عدا المطوّرين"],
            ["المعدات", "المدير العام، التطوير، العمليات"],
            ["المورّدون", "المدير العام، الأعمال، العمليات"],
            ["المالية", "المدير العام، الأعمال، الحسابات، رئيس القسم التقني"],
            ["المستخدمون", "المدير العام ورؤساء الأقسام"],
            ["النشاط والأتمتة", "المدير العام فقط"],
        ],
        rtl=rtl,
        col_widths=[4, 6],
    )
    note_box(doc, "المطوّر لا يرى موعد التركيب الحقيقي ولا أي أرقام مالية.", rtl=rtl)

    heading(doc, "5.  دورة حياة المشروع", rtl=rtl)
    p(doc, "المسار المعتاد. الزر يظهر فقط إذا كان دورك يملك هذه الخطوة.", rtl=rtl)
    add_table(
        doc,
        ["الخطوة", "الحالة", "من يعمل", "ماذا يفعل"],
        [
            ["1", "تم تقديم الملخص", "رئيس القسم التقني", "يكلّف الفريق ويحدد موعد التطوير."],
            ["2", "تم التكليف", "المطوّر", "يبدأ التطوير."],
            ["3", "قيد التطوير", "المطوّر", "يرسل العمل للمراجعة الداخلية."],
            ["4", "مراجعة داخلية", "رئيس القسم التقني", "يعتمد فقط بعد وجود الدليل وفيديو العرض."],
            ["5", "جاهز للتركيب", "العمليات", "تبدأ التركيب في الموقع."],
            ["6", "جارٍ التركيب", "العمليات", "تعلّم أنه تم التركيب وتكتب التقرير."],
            ["7", "تم التركيب", "الأعمال", "تعلّم المشروع مكتملًا."],
            ["8", "مكتمل", "—", "التقرير الشهري يستخدم هذا التاريخ."],
        ],
        rtl=rtl,
        col_widths=[1.0, 2.4, 2.2, 4.4],
    )
    bullets(
        doc,
        [
            [
                ("رئيس القسم التقني", True),
                " يستطيع إرجاع المشروع ",
                ("إلى التطوير", True),
                " من المراجعة الداخلية أو من «جاهز للتركيب».",
            ],
            ["العمليات تستطيع إيقاف التركيب مؤقتًا (العودة إلى جاهز للتركيب)."],
            "قد يكون المشروع متوقفًا مؤقتًا أو ملغى في البيانات؛ العمل اليومي يتبع المسار أعلاه.",
        ],
        rtl=rtl,
    )
    heading(doc, "من يفعل ماذا", 2, rtl=rtl)
    numbered(
        doc,
        [
            [
                ("الأعمال", True),
                " ينشئ الملخص (الاسم، العميل، الوصف، المراجع، التصميم، المحتوى، ملاحظات أخرى). يضع الميزانية و",
                ("موعد التركيب الحقيقي", True),
                " عند معرفته. يمكن إضافة مواعيد عرض العمل الجاري للعميل.",
            ],
            [
                ("رئيس القسم التقني", True),
                " يفتح ",
                ("التكليفات", True),
                "، يختار المطوّرين، يحدد موعد التطوير والمراحل.",
            ],
            [
                ("المطوّرون", True),
                " يبدأون التطوير، ينهون المراحل، يكتبون ",
                ("دليل التركيب", True),
                " (إنجليزي و/أو عربي) ويضعون رابط ",
                ("فيديو العرض", True),
                " المصوَّر على نفس التجهيز الذي سيُركَّب.",
            ],
            "لا يمكن اعتماد «جاهز للتركيب» قبل الخطوات وفيديو العرض.",
            [
                ("العمليات", True),
                " تستخدم ",
                ("التركيبات", True),
                "، تعيّن فنيًا مسؤولًا، تبدأ العمل، تكتب التقرير، تعلّم أنه تم التركيب.",
            ],
            [
                ("الأعمال", True),
                " تعلّم المشروع ",
                ("مكتملًا", True),
                " بعد التركيب (هذا التاريخ يدخل التقرير الشهري).",
            ],
            [("الحسابات", True), " تسجّل التحصيل والمصروفات في أي وقت بعد وجود الميزانية."],
        ],
        rtl=rtl,
    )

    heading(doc, "6.  فتح المشروع — التبويبات", rtl=rtl)
    p(doc, "المشروع مقسوم إلى تبويبات. يمكنك فتح أي تبويب مسموح لدورك.", rtl=rtl)
    p(doc, "إذا فتحت المشروع من الرئيسية أو المشاريع، التبويب الأول حسب الدور:", rtl=rtl, bold=True, space_after=6)
    add_table(
        doc,
        ["أنت", "يفتح على"],
        [
            ["محاسب", "المالية"],
            ["عمليات", "التركيب"],
            ["مطوّر / رئيس القسم التقني", "التطوير"],
            ["أعمال / المدير العام", "الملخص"],
        ],
        rtl=rtl,
        col_widths=[5, 5],
    )
    p(doc, "إذا ضغطت المشروع من صفحة أخرى، تلك الصفحة هي التي تحدّد التبويب:", rtl=rtl, bold=True, space_after=6)
    add_table(
        doc,
        ["ضغطت من", "يفتح"],
        [
            ["المالية", "المالية"],
            ["التركيبات", "التركيب"],
            ["التكليفات", "التطوير"],
            ["الطلبات", "الطلبات"],
            ["المعدات", "المعدات"],
        ],
        rtl=rtl,
        col_widths=[5, 5],
    )

    heading(doc, "7.  الرئيسية", rtl=rtl)
    p(
        doc,
        "ترحيب، دورك، وإن كنت ترى كل المشاريع: أعداد المسار ومواعيد التطوير القريبة. اللوحات تظهر فقط إن كان فيها شيء لك:",
        rtl=rtl,
    )
    bullets(
        doc,
        [
            [("مشاريعي", True), " — الأعمال المكلَّف بها"],
            [("التركيبات القادمة", True), " — إن كنت من العمليات"],
            [("ملاحظات العميل", True), " المسندة إليك، أو بانتظار الأعمال"],
            [("حالة النظام", True), " — المدير العام فقط"],
        ],
        rtl=rtl,
    )

    heading(doc, "8.  المشاريع", rtl=rtl)
    bullets(
        doc,
        [
            [("المطوّر", True), " يرى مشاريعه فقط."],
            "الباقون يرون كل المشاريع (ما عدا الملغى في بعض القوائم).",
            [("ملخص جديد", True), " لقسم الأعمال (والمدير العام)."],
            [
                ("تحميل Excel", True),
                " للمدير العام ورئيسي الأعمال والحسابات. يمكن تجميع الأرباع حسب تاريخ الملخص أو تاريخ التركيب.",
            ],
        ],
        rtl=rtl,
    )
    p(
        doc,
        "كل بطاقة تعرض الحالة والعميل والفريق وموعد التطوير، وإن سُمح لك: موعد التركيب والميزانية.",
        rtl=rtl,
    )

    heading(doc, "9.  تبويبات المشروع بالتفصيل", rtl=rtl)

    heading(doc, "الملخص", 2, rtl=rtl)
    p(doc, "ما طلبه العميل. الأعمال تعدّله. مواعيد عرض العمل الجاري هنا أيضًا.", rtl=rtl)

    heading(doc, "الملاحظات", 2, rtl=rtl)
    p(
        doc,
        "سجّل ما قاله العميل (واتساب، اتصال، بريد، اجتماع، زيارة موقع، أو رابط العميل). تُقسَّم إلى بنود: تصميم، محتوى، تقني، طلب جديد، عام.",
        rtl=rtl,
    )
    p(doc, "المسار المعتاد:", rtl=rtl, bold=True, space_after=4)
    numbered(
        doc,
        [
            ["الأعمال أو العمليات ", ("يسجّلان", True), " الملاحظة (يمكن تقسيمها تلقائيًا إلى بنود)."],
            [
                "إن كان ",
                ("طلبًا جديدًا خارج الملخص", True),
                " ينتظر قرار الأعمال: ضمن المشروع / تكلفة إضافية / مرفوض.",
            ],
            ["رئيس القسم التقني أو الأعمال ", ("يسندان", True), " البند."],
            ["المسؤول (أو مطوّرو المشروع) يعلّمون ", ("قيد التنفيذ", True), " ثم ", ("منجز", True), "."],
            [
                "المطوّر قد يطلب ",
                ("أيامًا إضافية", True),
                " على موعد التطوير؛ رئيس القسم التقني يقبل أو يرفض.",
            ],
            [
                "الأعمال يؤكد أن ",
                ("العميل اعتمد", True),
                "، أو العميل يفعل ذلك من رابطه.",
            ],
        ],
        rtl=rtl,
    )
    note_box(
        doc,
        "رابط العميل: الأعمال يشغّل رابطًا خاصًا. العميل يرى حالة مبسطة (بدون خطوات داخلية)، يضيف تعليقًا، ويقول إن البند لم يصبح صحيحًا بعد. نقاش الفريق يبقى داخليًا. البنود العاجلة ظاهرة بوضوح.",
        rtl=rtl,
    )

    heading(doc, "الملفات", 2, rtl=rtl)
    p(
        doc,
        "رفع مشترك (سحابة): الملخص، ملفات العميل، الهوية، التصميم، المحتوى، المعاينة، التسليم. حد 100 ميجابايت للملف.",
        rtl=rtl,
    )
    p(
        doc,
        "ملفات العمل الثقيلة (Unreal وC4D وPremiere) تبقى على قرص المكتب في المجلدات الظاهرة في الصفحة — لا تُرفع.",
        rtl=rtl,
    )
    p(
        doc,
        "المطوّر يدير ملفات مشاريعه فقط. الأعمال والعمليات ورئيس القسم التقني والمدير العام أوسع صلاحية.",
        rtl=rtl,
    )

    heading(doc, "التطوير", 2, rtl=rtl)
    p(doc, "الفريق، موعد التطوير، المراحل، أزرار الحالة، ودليل التركيب.", rtl=rtl)
    bullets(
        doc,
        [
            [("رئيس القسم التقني:", True), " تكليف / إزالة مطوّرين، الموعد، إضافة مراحل، الاعتماد أو الإرجاع."],
            [
                ("مطوّرو المشروع:", True),
                " بدء العمل، الإرسال للمراجعة، تعليم المراحل، كتابة الدليل (أي لغة؛ الأخرى تُترجم عند الحفظ إن كانت فارغة).",
            ],
        ],
        rtl=rtl,
    )

    heading(doc, "الطلبات", 2, rtl=rtl)
    p(
        doc,
        "داخل المشروع وعلى صفحة الطلبات. اطلب من شخص ما يحتاجه المشروع: متوقف، يحتاج بيانات، يحتاج ملفات، أو غير ذلك. يظهر الطلب فقط للمعنيين.",
        rtl=rtl,
    )
    bullets(
        doc,
        [
            ["المستلم (أو رئيس قسمه) ينهي ", ("جزأه", True), "."],
            "عند إنهاء كل الأجزاء يُغلق الطلب.",
            ["من رفع الطلب (أو المدير العام) يستطيع ", ("الإلغاء", True), "."],
            ["الرقم على ", ("الطلبات", True), " في القائمة هو ما ينتظرك."],
        ],
        rtl=rtl,
    )

    heading(doc, "المعدات (على المشروع)", 2, rtl=rtl)
    p(
        doc,
        "اطلب وحدات لفترة زمنية. إن وُجدت وحدات حرة يُعتمد الطلب فورًا. وإلا يدخل قائمة الانتظار حتى يتوفر المخزون.",
        rtl=rtl,
    )
    p(
        doc,
        "بعد انتهاء الاستخدام علّم تم الإرجاع (أو إلغاء). المدير العام ورئيسا التطوير والعمليات يحدّثان الكتالوج (الكمية والحالة: نشط / صيانة / معطل / خارج الخدمة).",
        rtl=rtl,
    )

    heading(doc, "التركيب", 2, rtl=rtl)
    p(doc, "مخفي عن المطوّرين.", rtl=rtl)
    bullets(
        doc,
        [
            [("الأعمال", True), " يحدد الموعد الحقيقي ورابط الموقع والملاحظات."],
            [("العمليات", True), " تعيّن الفني المسؤول، تبدأ التركيب، توقفه، تعلّم أنه تم، تكتب التقرير."],
            [("الأعمال", True), " تعلّم ", ("مكتمل", True), " بعد التركيب."],
        ],
        rtl=rtl,
    )

    heading(doc, "المالية", 2, rtl=rtl)
    p(doc, "مخفي عن المطوّرين والعمليات.", rtl=rtl)
    add_table(
        doc,
        ["الرقم", "المعنى"],
        [
            ["الميزانية", "المبلغ المتفق عليه. لا تنقص عند إضافة مصروف."],
            ["الإيراد", "ما حُصّل. يبقى صفرًا حتى تسجيل دفعة."],
            ["المتبقي", "الميزانية ناقص الإيراد (ما لم يُحصَّل بعد)."],
            ["المصروفات", "التكاليف."],
            ["الربح", "الإيراد ناقص المصروفات. يظهر بالأحمر إن زادت التكاليف على التحصيل."],
        ],
        rtl=rtl,
        col_widths=[2.4, 7.6],
    )
    p(doc, "الصلاحيات:", rtl=rtl, bold=True, space_after=4)
    bullets(
        doc,
        [
            [("عرض:", True), " المدير العام، الأعمال، الحسابات، رئيس القسم التقني"],
            [("تعديل الميزانية:", True), " المدير العام، الأعمال، الحسابات"],
            [("تسجيل الدفعات والمصروفات:", True), " المدير العام والحسابات فقط"],
        ],
        rtl=rtl,
    )
    note_box(
        doc,
        "كيف تسجّل الحسابات عدة قيود: افتح تسجيل دفعة أو مصروف ← املأ سطرًا ← إضافة (كرر) ← حفظ مرة واحدة. يمكن إزالة سطر من القائمة قبل الحفظ. إن كان آخر سطر مملوءًا وضغطت حفظ بدون إضافة، يُحفظ معه.",
        rtl=rtl,
    )
    p(doc, "صفحة المالية تعرض نفس الأرقام لكل مشروع.", rtl=rtl)

    heading(doc, "10.  التكليفات (رئيس القسم التقني)", rtl=rtl)
    p(
        doc,
        "الملخصات المنتظرة كبطاقات، ثم كل مطوّر ومشاريعه النشطة تحته (أقرب موعد أولًا). افتح البطاقة لتكليف الأشخاص والموعد والمراحل في تبويب التطوير.",
        rtl=rtl,
    )

    heading(doc, "11.  تقويم التركيبات", rtl=rtl)
    p(
        doc,
        "عرض شهري لمواعيد الموقع، «جارٍ التركيب الآن»، ومشاريع بلا موعد. الأعمال يضع الموعد على المشروع. العمليات تنفّذ اليوم من تبويب التركيب.",
        rtl=rtl,
    )

    heading(doc, "12.  كتالوج المعدات", rtl=rtl)
    p(
        doc,
        "المخزون، الحر مقابل المستخدم، أي مشروع يحجز وحدات، قائمة الانتظار. الطلب من المشروع. رئيسا التطوير والعمليات (والمدير العام) يضيفون ويعدّلون الأصناف.",
        rtl=rtl,
    )

    heading(doc, "13.  المورّدون", rtl=rtl)
    p(doc, "اسم الشركة، جهة الاتصال، الهاتف، العنوان، والمعدات المستأجرة. للأعمال والعمليات والمدير العام.", rtl=rtl)

    heading(doc, "14.  المستخدمون", rtl=rtl)
    bullets(
        doc,
        [
            "المدير العام يرى الجميع ويعيّن الرؤساء.",
            [
                "رئيس القسم يرى قسمه فقط ويضيف أعضاء ",
                ("بكلمة مرور مؤقتة", True),
                " (8 أحرف على الأقل). العضو الجديد يغيّرها في أول دخول.",
            ],
            [
                "إعادة تعيين كلمة المرور أو التعطيل أو التفعيل من صفحة الشخص. لا تعطّل نفسك؛ غيّر كلمة مرورك من ",
                ("الحساب", True),
                ".",
            ],
        ],
        rtl=rtl,
    )

    heading(doc, "15.  رسائل البريد التي يرسلها النظام", rtl=rtl)
    p(doc, "لا تحتاج تشغيلها يدويًا. أمثلة:", rtl=rtl)
    bullets(
        doc,
        [
            "تم تكليفك بمشروع",
            "طلب يحتاج جزأك",
            "أُلغي طلب",
            "ملاحظة عميل مسندة إليك",
            "موعد التطوير بعد 3 أيام أو غدًا، أو متأخر",
            "التركيب بعد يومين، أو متأخر",
            "معدات ما زالت محجوزة",
            [
                ("ملخص أسبوعي", True),
                " (الرؤساء والمدير العام): المسار، التركيبات القادمة، المواعيد المتأخرة، الطلبات المفتوحة، المعدات المستخدمة",
            ],
            [
                ("تقرير شهري", True),
                " (رئيس الأعمال والمدير العام): المشاريع المكتملة في الشهر السابق مع ملف Excel",
            ],
        ],
        rtl=rtl,
    )

    heading(doc, "16.  للمدير العام فقط", rtl=rtl)
    bullets(
        doc,
        [
            [("النشاط", True), " — من غيّر ماذا"],
            [("الأتمتة", True), " — صحة البريد، وتشغيل المهام اليومية / الأسبوعية / الشهرية عند الحاجة"],
        ],
        rtl=rtl,
    )

    heading(doc, "17.  بطاقات سريعة", rtl=rtl)

    heading(doc, "أنا من الأعمال", 2, rtl=rtl)
    numbered(
        doc,
        [
            [("المشاريع ← ملخص جديد", True)],
            ["ضع ", ("الميزانية", True), " و", ("موعد التركيب", True), " عندما تعرفهما"],
            ["سجّل ", ("ملاحظات العميل", True), "؛ احسم نطاق «الطلب الجديد»"],
            ["شغّل ", ("رابط العميل", True), " عندما يرد بنفسه"],
            ["بعد أن تعلّم العمليات أنه تم التركيب ← ", ("تعليم مكتمل", True)],
        ],
        rtl=rtl,
    )

    heading(doc, "أنا رئيس القسم التقني", 2, rtl=rtl)
    numbered(
        doc,
        [
            [("التكليفات", True), " — التقط الملخصات المنتظرة"],
            "كلّف المطوّرين والموعد والمراحل",
            [
                "بعد إرسال المراجعة ← راجع الدليل والفيديو ← ",
                ("اعتماد: جاهز للتركيب", True),
                " (أو أرجع)",
            ],
            "اقبل أو ارفض أيام التطوير الإضافية من بنود الملاحظات",
        ],
        rtl=rtl,
    )

    heading(doc, "أنا مطوّر", 2, rtl=rtl)
    numbered(
        doc,
        [
            [("الرئيسية / المشاريع", True), " — أعمالك فقط"],
            [("بدء التطوير", True), " ← المراحل ← اكتب الدليل وفيديو العرض"],
            [("إرسال للمراجعة الداخلية", True)],
            "اطلب معدات أو ارفع طلبًا إن توقفت",
            "لن ترى المال ولا موعد التركيب الحقيقي",
        ],
        rtl=rtl,
    )

    heading(doc, "أنا من العمليات", 2, rtl=rtl)
    numbered(
        doc,
        [
            ["تقويم ", ("التركيبات", True)],
            [
                "في تبويب ",
                ("التركيب", True),
                ": الفني المسؤول، ",
                ("بدء التركيب", True),
                "، التقرير، ",
                ("تعليم تم التركيب", True),
            ],
            "المعدات والمورّدون حسب الحاجة",
            "يمكنك تسجيل ملاحظة عميل من الموقع؛ الأعمال تحسم النطاق الإضافي",
        ],
        rtl=rtl,
    )

    heading(doc, "أنا من الحسابات", 2, rtl=rtl)
    numbered(
        doc,
        [
            [("المالية", True), " أو تبويب ", ("المالية", True), " في المشروع (يفتح تلقائيًا)"],
            ["راجع ", ("الميزانية", True)],
            [("أضف", True), " الدفعات والمصروفات للقائمة ثم ", ("احفظ", True)],
            ["راقب ", ("الإيراد", True), " (التحصيل) و", ("المتبقي", True), " و", ("الربح", True)],
        ],
        rtl=rtl,
    )

    heading(doc, "18.  إن نقص شيء", rtl=rtl)
    add_table(
        doc,
        ["ما تراه", "المعنى المعتاد"],
        [
            ["عنصر قائمة غير ظاهر", "دورك غير مخصّص له. هذا طبيعي."],
            ["لا تفتح مشروعًا", "أنت مطوّر ولست في الفريق. اطلب من رئيس القسم التقني."],
            ["لا تغيّر الحالة", "الخطوة السابقة لم تتم، أو الدليل/الفيديو ناقص."],
            ["فشل الدخول", "كلمة المرور، أو الحساب معطّل. راجع رئيسك."],
        ],
        rtl=rtl,
        col_widths=[3.5, 6.5],
    )
    p(doc, "أسئلة سياسة العمل في الشركة (وليس الأزرار) تُوجَّه لرئيس قسمك.", rtl=rtl)

    doc.save(path)


def main():
    en = OUT / "ARTECH-Studio-OS-Handbook-EN.docx"
    ar = OUT / "ARTECH-Studio-OS-Handbook-AR.docx"
    build_english(en)
    build_arabic(ar)
    print(en)
    print(ar)


if __name__ == "__main__":
    main()
